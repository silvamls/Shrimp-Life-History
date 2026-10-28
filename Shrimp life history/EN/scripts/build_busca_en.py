#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_busca_en.py  -  English version of the search workbook ("busca_2026_lista_referencias.xlsx")
=====================================================================================================
INPUT : busca_v6.xlsx   (sheets: referencias, log_de_busca, resumo, candidatos_busca_web)
OUTPUT: busca_en.xlsx   (sheets: records, search_log, summary, web_search_candidates)

HOW IT WORKS
  1. Free text was translated beforehand (three parallel translation passes -> en/b1_en.json, b2_en.json,
     b3_en.json), keyed by 'sheet|column_index|unique_string_index' (index into en/busca_strings.json).
     This script only APPLIES those translations (function tr()).
  2. Column 12 of 'referencias' ("Closest record in master") has the form
        "[sim. 0.75] 'Rodrigues et al., 1992' (linha 309): <bibliographic reference>"
     and is translated by regular expression (similarity -> similarity, linha -> row); the reference
     itself is kept as published.
  3. Working-note wording ("the user", "Matheus") is rewritten for a publication context (TEXT_FIXES).
  4. Three rows are added at the top of the search log with the real search timeline given by the authors
     (month-level dates) and the two rows that said 'subscription databases NOT ACCESSED' are re-labelled,
     because those databases were searched manually by the authors; the log rows document only the
     AI-assisted web-retrieval step.
  5. Titles, authors, DOIs/URLs and bibliographic strings are never translated.
"""
import json, re, copy
import openpyxl
from openpyxl.styles import Alignment

SRC, DST, D = '/home/claude/org/busca_v6.xlsx', '/home/claude/org/busca_en.xlsx', '/home/claude/org/en/'
SHEETS = {'referencias': 'records', 'log_de_busca': 'search_log', 'resumo': 'summary',
          'candidatos_busca_web': 'web_search_candidates'}

strings = json.load(open(D + 'busca_strings.json'))
EN = {}
for f in ('b1', 'b2', 'b3'):
    EN.update(json.load(open(D + f + '_en.json')))

# wording for a publication context ------------------------------------------------------------
TEXT_FIXES = [
    ("on the user's instruction", "by decision of the authors"),
    ("on the user's instructions", "by decision of the authors"),
    ("User correction", "Author correction"), ("user correction", "author correction"),
    ("User remarks", "Author remarks"), ("user remarks", "author remarks"),
    ("the user's", "the authors'"), ("The user's", "The authors'"),
    ("Matheus (13 Oct 2026)", "M.S.S. (13 Oct 2026)"), ("message from Matheus", "message from the author (M.S.S.)"),
    ("reported by Matheus", "reported by the author"),
    ("I did not find a numerical CL50", "no numerical CL50 was found"),
    ("the user", "the authors"),
    ("see sheet 'notas'", "see sheet 'notes'"),
]
def fixes(t):
    for a, b in TEXT_FIXES: t = t.replace(a, b)
    return t

def tr(sheet, col0, v):
    lst = strings[f'{sheet}|{col0}']
    return fixes(EN[f'{sheet}|{col0}|{lst.index(v)}'])

def col12(v):
    """'[sim. 0.75] 'X' (linha 309): ref' -> '[similarity 0.75] 'X' (row 309): ref'"""
    v = re.sub(r'^\[sim\. ([\d.]+)\]', r'[similarity \1]', v)
    v = re.sub(r'\(linha (\d+)\)', r'(row \1)', v)
    v = re.sub(r'\(linhas ([\d,\s–-]+)\)', r'(rows \1)', v)
    return v

wb = openpyxl.load_workbook(SRC)

# ---- translate every string cell ------------------------------------------------------------------
HDR12 = 'Closest record in master (for checking)'
for ws in wb:
    old = ws.title
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str): continue
            col0 = c.column - 1
            key = f'{old}|{col0}'
            if old == 'referencias' and col0 == 12:
                c.value = HDR12 if c.row == 1 else col12(v); continue
            if old == 'referencias' and col0 in (1, 2, 4) and c.row > 1: continue    # authors / titles / DOI: keep
            if old == 'candidatos_busca_web' and col0 in (1, 2, 4) and c.row > 1: continue
            if key in strings and v in strings[key]:
                c.value = tr(old, col0, v)
            elif not (old == 'referencias' and col0 in (1, 2, 4)):
                print('NOT TRANSLATED', old, c.coordinate, v[:60])

# ---- search log: re-label the two 'not accessed/not performed' rows and add the real timeline -----
ws = wb['log_de_busca']
for r in range(2, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if isinstance(a, str) and a.startswith('Web of Science, Scopus, ASFA'):
        ws.cell(r, 2).value = 'subscription databases (institutional access)'
        ws.cell(r, 3).value = 'NOT ACCESSED BY THE ASSISTED WEB-RETRIEVAL STEP'
        ws.cell(r, 4).value = ('These databases were searched manually by the authors (WoS, Scopus, ASFA and the CAPES Periodicals '
                               'Portal; see the first rows of this log). Zoological Record and LILACS were not searched.')
        ws.cell(r, 1).value = 'Web of Science, Scopus, ASFA, Zoological Record, LILACS (assisted web-retrieval step only)'
    if isinstance(a, str) and a.startswith('Google Scholar (first 200'):
        ws.cell(r, 4).value = 'Foreseen in the protocol; not performed (see the manual citation searches in the first rows of this log).'

timeline = [
 ('Bibliographic databases: Web of Science, Scopus, ASFA, CAPES Periodicals Portal',
  'manual searches with the same search strings in all databases; two searchers, results checked independently',
  'COMPLETED — Jul 2026 (A.C.A.F.), Aug 2026 (M.S.S.), early Oct 2026 (M.S.S., final update)',
  'Search sessions of about three per week in July and in August (A.C.A.F. covered all sources in July; M.S.S. in August); '
  'two final update sessions in early October 2026 to retrieve remaining records. No manual database or citation searches in September 2026.'),
 ('Other methods: backward and forward citation searching, grey literature, complementary sources',
  'manual; citation chasing from included reports, repositories and institutional portals',
  'COMPLETED — Jul 2026 (A.C.A.F.), Aug 2026 (M.S.S.), early Oct 2026 (M.S.S., final update)',
  'Same schedule as the database searches (mainly direct and reverse citation searching). Entry dates in the master spreadsheet and the '
  'creation dates of the individual per-article spreadsheets document the collection dates.'),
 ('Independent comparison of the searches (second searcher)',
  'all sources searched by both authors and compared record by record',
  'COMPLETED — Jul–Oct 2026',
  'Divergences (1–2 sources) arose from internal inconsistencies in the published reports and were resolved by re-checking the article. '
  'A.C.A.F. = Anna Carolina Arnaldo de Figueiredo (ORCID 0000-0003-2569-972X); M.S.S. = lead searcher/data extractor.'),
]
# insert the 3 rows right below the header, copying the style of an existing body row
ws.insert_rows(2, amount=len(timeline))
for i, vals in enumerate(timeline):
    r = 2 + i
    for j, v in enumerate(vals, start=1):
        c = ws.cell(r, j); c.value = v
        src = ws.cell(2 + len(timeline), j)                       # first original row (now shifted down)
        c.font = copy.copy(src.font); c.fill = copy.copy(src.fill); c.border = copy.copy(src.border)
        c.alignment = Alignment(wrap_text=True, vertical='top')
# the description header of the log
ws.cell(1, 1).value = 'Source consulted'; ws.cell(1, 2).value = 'Access mode'
ws.cell(1, 3).value = 'Status'; ws.cell(1, 4).value = 'Note'

# ---- sheet names, view -------------------------------------------------------------------------------
for ws in wb: ws.title = SHEETS[ws.title]
wb.active = 0
wb.save(DST)
print('saved', DST)
